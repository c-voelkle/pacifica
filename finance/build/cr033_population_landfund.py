"""Apply the population change list (P1-P6), the CR-033 Land Fund and a larger
Events table to the financial workbook.

    python3 finance/build/cr033_population_landfund.py IN.xlsx OUT.xlsx [--rates rates.csv]

`rates.csv` (year, rate) supplies Paths C, the natural-increase rate from the
population model (population/changes.py, column paths_C_model). Without it,
Paths C is left as it is. The script never recalculates; run LibreOffice
(xlsx skill scripts/recalc.py) afterwards. Every new number is a labelled
input with its source in a cell comment.
"""
from __future__ import annotations

import argparse
import copy
import csv
import math
import re
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.defined_name import DefinedName

FIRST_ROW, LAST_ROW = 6, 92            # 1846 ... 1932 on every year sheet
EVENTS_LAST = 400                      # was 174
BLUE = Font(name="Arial", size=10, color="FF0000FF")
BLACK = Font(name="Arial", size=10, color="FF000000")
GREEN = Font(name="Arial", size=10, color="FF008000")
BOLD = Font(name="Arial", size=10, bold=True)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
AUTHOR = "CR-033 / population model, 27 Oct 2026"


def year_row(y: int) -> int:
    return y - 1840


def note(cell, text: str) -> None:
    cell.comment = Comment(text, AUTHOR, width=320, height=160)


# ---------------------------------------------------------------------------
# 1. Events table: room for rows up to EVENTS_LAST
# ---------------------------------------------------------------------------

def expand_events(wb) -> int:
    pat = re.compile(r"(Events!\$[A-Z]+\$6:\$[A-Z]+\$)174\b")
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "Events!" in c.value and pat.search(c.value):
                    c.value = pat.sub(lambda m: f"{m.group(1)}{EVENTS_LAST}", c.value)
                    n += 1
    E = wb["Events"]
    src = 174
    for r in range(src + 1, EVENTS_LAST + 1):
        for col in range(1, 14):
            s, d = E.cell(src, col), E.cell(r, col)
            if s.has_style:
                d._style = copy.copy(s._style)
        E.cell(r, 9).value = f'=IF(H{r}="","",H{r}*(C{r}-B{r}+1))'
    E["A4"] = (f"Rows 6-{EVENTS_LAST} are read by the model (widened from 174 on 27 Oct 2026). "
               "Add new events in the first empty row.")
    E["A4"].font = Font(name="Arial", size=10, italic=True)
    return n


def add_event(wb, values: list) -> int:
    E = wb["Events"]
    r = next(r for r in range(6, EVENTS_LAST + 1) if E.cell(r, 1).value in (None, ""))
    last_id = max(int(str(E.cell(x, 1).value)[1:]) for x in range(6, r) if str(E.cell(x, 1).value).startswith("E"))
    E.cell(r, 1).value = f"E{last_id + 1}"
    for i, v in enumerate(values, start=2):
        if i == 9:
            continue
        E.cell(r, i).value = v
    return r


# ---------------------------------------------------------------------------
# Levers
# ---------------------------------------------------------------------------

def add_levers(wb, rows: list[tuple]) -> dict[str, str]:
    """rows: (name, label, story, stress, unit, status, source) or a header string."""
    L = wb["Levers"]
    r = L.max_row + 1
    style_src = {c: L.cell(275, c)._style for c in range(1, 12)}
    head_src = {c: L.cell(109, c)._style for c in range(1, 12)}
    refs = {}
    for item in rows:
        if isinstance(item, str):
            for c in range(1, 12):
                L.cell(r, c)._style = copy.copy(head_src[c])
            L.cell(r, 2).value = item
            r += 1
            continue
        name, label, story, stress, unit, status, source = item
        for c in range(1, 12):
            L.cell(r, c)._style = copy.copy(style_src[c])
        L.cell(r, 2).value = label
        L.cell(r, 3).value = story
        L.cell(r, 4).value = stress
        L.cell(r, 5).value = f"=IF(case_sel=1,C{r},D{r})"
        L.cell(r, 6).value = unit
        L.cell(r, 7).value = status
        note(L.cell(r, 3), source)
        if name in wb.defined_names:
            raise SystemExit(f"defined name {name!r} already exists ({wb.defined_names[name].attr_text})")
        wb.defined_names[name] = DefinedName(name, attr_text=f"Levers!$E${r}")
        refs[name] = f"Levers!$E${r}"
        r += 1
    return refs


# ---------------------------------------------------------------------------
# 2. Population (P1-P6)
# ---------------------------------------------------------------------------

def apply_population(wb, rates: dict[int, float] | None) -> None:
    add_levers(wb, [
        "— CR-033 and the population model (27 Oct 2026) —",
        ("dem01_mult", "DEM-01: arrivals from the Americas, Europe and China, 1848-1860 (× sheet path)", 2.10, 2.10,
         "×", "Approved",
         "Author decision 27 Oct 2026 (DEM-01, about 1.2m in 1861); population/data/assumptions.toml [dem01]. "
         "Counterfactuals: light licences, assisted passage, secure titles (CR-033 A1)."),
        ("dem01_target", "DEM-01: residents at the start of 1861 (check)", 1200000, 1200000, "persons", "Approved",
         "war_base_assumptions.md DEM-01; checked within 10% on Checks!AK."),
        ("xo_year", "Extra arrivals by origin: second period starts", 1914, 1914, "year", "Approved",
         "Author 27 Oct 2026: Filipino migration mostly after 1914; population/data/assumptions.toml [extra_origin]."),
        ("xo_jp1", "Extra arrivals from Japan, before that year", 0.20, 0.20, "share", "Approved",
         "Author 27 Oct 2026: Japanese predominance (CR-033 A11). Americas & Europe take the remainder."),
        ("xo_cn1", "Extra arrivals from China, before that year", 0.10, 0.10, "share", "Approved", "As above."),
        ("xo_ph1", "Extra arrivals from the Philippines, before that year", 0.0, 0.0, "share", "Approved", "As above."),
        ("xo_jp2", "Extra arrivals from Japan, from that year", 0.32, 0.32, "share", "Approved", "As above."),
        ("xo_cn2", "Extra arrivals from China, from that year", 0.13, 0.13, "share", "Approved", "As above."),
        ("xo_ph2", "Extra arrivals from the Philippines, from that year", 0.05, 0.05, "share", "Approved", "As above."),
    ])
    L = wb["Levers"]
    note(L["C8"], "Includes about 150,000 Native Californians (author decision 27 Oct 2026; population/README.md decision 2).")

    P = wb["People"]
    for r in range(FIRST_ROW, LAST_ROW + 1):
        for col in "EF":
            f = P[f"{col}{r}"].value
            if "dem01_mult" not in f:
                P[f"{col}{r}"].value = f + "*IF(AND($A{r}>=1848,$A{r}<=1860),dem01_mult,1)".format(r=r)
        if r > FIRST_ROW:
            share = lambda a, b: f"Labour!AL{r}*IF($A{r}>=xo_year,{b},{a})"
            P[f"N{r}"].value = f"=N{r-1}*(1+Paths!C{r})+F{r}*(1-dep_asia)+" + share("xo_cn1", "xo_cn2")
            P[f"O{r}"].value = f"=O{r-1}*(1+Paths!C{r})+G{r}*(1-dep_asia)+" + share("xo_jp1", "xo_jp2")
            P[f"P{r}"].value = f"=P{r-1}*(1+Paths!C{r})+H{r}*(1-dep_asia)+" + share("xo_ph1", "xo_ph2")
            P[f"Q{r}"].value = "=0"
            P[f"R{r}"].value = "=0"
    note(P["N5"], "From 27 Oct 2026 N, O and P also take their origin share of the extra arrivals (Labour AL; "
                  "Levers xo_*). Q and R are folded in and kept at zero, so nothing is counted twice.")
    P["Q5"].value = "Settled guest workers of Asian origin (now in N-P)"
    P["R5"].value = "Asian talent migrants and families (now in N-P)"

    # P4: memo columns Y, Z from the population model
    model = read_model_memo()
    for col, key, head in (("Y", "native_hawaiian_model", "Native Hawaiians (memo, population model)"),
                           ("Z", "native_californian_model", "Native Californians (memo, population model)")):
        P.column_dimensions[col].width = 16
        for r in (4, 5):
            P[f"{col}{r}"]._style = copy.copy(P[f"X{r}"]._style)
        P[f"{col}4"].value = "Memo (not in totals)" if col == "Y" else None
        P[f"{col}5"].value = head
        for r in range(FIRST_ROW, LAST_ROW + 1):
            c = P[f"{col}{r}"]
            c._style = copy.copy(P[f"X{r}"]._style)
            c.font = BLUE
            c.value = round(model.get(1840 + r, {}).get(key, 0.0))
        note(P[f"{col}5"], "End-of-year values from population/changes.py (outputs/workbook_changes.csv). "
                           "Memo only: the totals carry these groups through Paths C. MODEL.")

    # P1: natural-increase rate
    if rates:
        Pa = wb["Paths"]
        for r in range(FIRST_ROW, LAST_ROW + 1):
            y = 1840 + r
            if y in rates:
                Pa[f"C{r}"].value = round(rates[y], 6)
        note(Pa["C5"], "From 27 Oct 2026: the population model's implied rate (births minus deaths over residents), "
                       "by age, sex and group with canon health and fertility paths and the Native Hawaiian and "
                       "Native Californian paths (population/changes.py, column paths_C_model). Iterated with the "
                       "workbook until stable. MODEL.")

    # P5: checks
    C = wb["Checks"]
    for col, head in (("AJ", "Native Hawaiian low in 1893-97"), ("AK", "1861 residents near DEM-01")):
        for r in (4, 5):
            C[f"{col}{r}"]._style = copy.copy(C[f"AI{r}"]._style)
        C[f"{col}5"].value = head
    C["AJ4"].value = "Population (CR-033)"
    for r in range(FIRST_ROW, LAST_ROW + 1):
        C[f"AJ{r}"]._style = copy.copy(C[f"AI{r}"]._style)
        C[f"AK{r}"]._style = copy.copy(C[f"AI{r}"]._style)
        low = "MATCH(MIN(People!$Y$16:$Y$92),People!$Y$16:$Y$92,0)+1855"   # from 1856, when Hawaii joins
        C[f"AJ{r}"].value = f'=IF($A{r}<>1895,"n/a",IF(AND({low}>=1893,{low}<=1897),"OK","Check"))'
        C[f"AK{r}"].value = f'=IF($A{r}<>1861,"n/a",IF(ABS(People!B{r}/dem01_target-1)<=0.1,"OK","Check"))'
        s = C[f"S{r}"].value
        if "AJ" not in s:
            C[f"S{r}"].value = s + f'+COUNTIF(AJ{r}:AK{r},"Check")'


def read_model_memo() -> dict[int, dict]:
    path = Path(__file__).resolve().parents[2] / "population" / "outputs" / "workbook_changes.csv"
    out = {}
    with open(path, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            out[int(row["year"])] = {k: float(v) for k, v in row.items() if k not in ("year",) and v not in ("",)
                                     and not v.startswith("(")}
    return out


# ---------------------------------------------------------------------------
# 3. Land Fund (CR-033 A5-A7, A17-A23)
# ---------------------------------------------------------------------------

# California land in farms, our timeline (census). The Pacifican need for land is
# taken as the same farm acreage: the extra people go mostly to towns and industry.
FARMLAND = {1850: 3893985, 1860: 8730035, 1870: 11427105, 1880: 16593742, 1890: 21427293,
            1900: 28828151, 1910: 27931444, 1920: 29365667, 1925: 27516955, 1930: 30442581}
FARMLAND_SOURCE = ("US Census of Agriculture: 1920 bulletin, California (www2.census.gov/library/publications/decennial/"
                   "1920/bulletins/agriculture/agriculture-ca.pdf) for 1850-1920; 1959 Census of Agriculture, California, "
                   "Table 1 for 1925-1930. Interpolated log-linearly between censuses; held before 1850 and after 1930.")


def farmland(y: int) -> float:
    pts = sorted(FARMLAND.items())
    if y <= pts[0][0]:
        return pts[0][1]
    if y >= pts[-1][0]:
        return pts[-1][1]
    for (y0, v0), (y1, v1) in zip(pts, pts[1:]):
        if y0 <= y <= y1:
            return math.exp(math.log(v0) + (math.log(v1) - math.log(v0)) * (y - y0) / (y1 - y0))


def apply_land_fund(wb) -> None:
    add_levers(wb, [
        "— CR-033: Land Fund and Native land (27 Oct 2026) —",
        ("lfund_on", "Land Fund and treaty policy (1 = on)", 1, 1, "switch", "Proposed", "CR-033 A5-A7, A17-A23 (draft)."),
        ("lfund_start", "First year of treaty purchases", 1851, 1851, "year", "Proposed", "CR-033 A6: treaties of 1851-52."),
        ("lfund_area", "Land area of California", 100000000, 100000000, "acres", "Model estimate",
         "California's land area is about 155,800 sq mi, about 100m acres; the 1850 Treaty of Monterey line is close to it."),
        ("lfund_home", "Native homelands reserved", 7400000, 7400000, "acres", "Model estimate",
         "Our timeline's eighteen treaties of 1851-52 offered 7.4m acres of reservations "
         "(Wikipedia, 'California Indian Reservations and Cessions'). Canon boundaries are OPEN (CR-033 A9)."),
        ("lfund_grants", "Land in the old Mexican grants", 8800000, 8800000, "acres", "Model estimate",
         "Our timeline confirmed 604 of 813 claims under the 1851 Land Act (Wikipedia, 'California Land Act of 1851'); "
         "about 8.8m acres is the commonly cited confirmed total, to verify. Not bought: privately owned."),
        ("lfund_farm_scale", "Farm land needed (× our timeline's California farm acreage)", 1.0, 1.0, "×", "Model estimate",
         "Pacifica's extra people go mainly to towns and industry; farm land is limited by what California's land can carry."),
        ("lfund_lead", "Years the Republic buys ahead of settlement", 5, 5, "years", "Model estimate",
         "CR-033 A17: buys district by district, ahead of settlement."),
        ("lfund_price", "Price paid to the nations (value of the annuity goods)", 0.10, 0.25, "$ an acre", "Model estimate",
         "OPEN in canon (A9). Order of magnitude of what the US and New Zealand governments paid (cents an acre); "
         "the stress case pays 2.5 times as much."),
        ("lfund_term", "Annuity term", 20, 20, "years", "Model estimate", "CR-033 A17: annuities over a long term; length OPEN."),
        ("lfund_sale", "Sale price to settlers", 1.25, 1.00, "$ an acre", "Model estimate",
         "The US minimum price for public land, $1.25 an acre (Land Act of 1820), is the price across the border. "
         "Stress case $1.00."),
        ("lfund_inst", "Instalment years for settlers", 5, 5, "years", "Model estimate", "CR-033 A17: settlers buy on instalments."),
        ("lfund_default", "Share of instalments never paid", 0.10, 0.20, "share", "Model estimate", "Model estimate."),
        ("lfund_survey", "Survey and register cost", 0.03, 0.05, "$ an acre", "Model estimate",
         "Rectangular survey under contract, a few cents an acre; to verify against GLO contract rates. "
         "Register running costs are covered by fees (A20) and left out."),
        ("lfund_grant_fee", "Share of grant land paid for the patent", 0.05, 0.05, "share", "Model estimate",
         "CR-033 A21: a grantee pays for a fast patent in land, which goes to the Land Fund."),
        ("lfund_grant_years", "Years to patent the old grants", 6, 8, "years", "Model estimate",
         "CR-033 A21: a short deadline; patents within months of survey."),
        ("lfund_rent", "Rent on leased homeland ground", 0.05, 0.05, "$ an acre a year", "Model estimate",
         "CR-033 A23: long leases at low rents."),
        ("lfund_lease_1932", "Share of homeland ground under lease by 1932", 0.30, 0.30, "share", "Model estimate",
         "CR-033 A23: land leaves Native hands through leases. Rises evenly from 1860."),
        ("lfund_fee", "Republic's fee on homeland rents", 0.10, 0.10, "share", "Model estimate", "CR-033 A22: the Republic keeps a fee."),
        ("lfund_police", "Commissioners and mounted police, 1852-1870", 300, 300, "men", "Model estimate",
         "CR-033 A7. Canada's North-West Mounted Police was authorised at 300 men in 1873."),
        ("lfund_police_late", "Commissioners and mounted police, 1871-1890", 150, 150, "men", "Model estimate",
         "Model estimate; county police take over after the 1887 Constitution."),
        ("lfund_police_cost", "Cost per man a year, at 1877 prices", 1150, 1400, "$", "Model estimate",
         "NWMP expenditure 1877-78 $345,959 (Tilley's budget speech, 3 March 1885) for about 300 men."),
        ("lfund_rate", "Interest on the Land Fund's balance with the Treasury", 0.05, 0.06, "a year", "Model estimate",
         "Model estimate, near the Republic's borrowing cost."),
    ])

    ws = wb.create_sheet("Land Fund", index=wb.sheetnames.index("Budget") + 1)
    ws.sheet_properties.tabColor = "7F6000"
    ws["A1"] = "Land Fund"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"] = ("CR-033 (draft): the Republic buys Native land by treaty for annuities in goods, sells it to settlers on "
                "instalments, takes grant land as the price of fast patents, and keeps a fee on homeland leases. "
                "Does it pay for itself? Inputs are on Levers (lfund_*); the farm-land path is our timeline's census.")
    ws["A3"] = ("Years run down the page. Blue = input · black = formula · green = link. $m in period money. "
                "Net cash to the Treasury feeds Budget column G. Police and commissioners are the wider policy's cost (A7).")
    for c in ("A2", "A3"):
        ws[c].font = Font(name="Arial", size=10, italic=True)
    heads = [
        ("A", "Year"),
        ("B", "Land in farms, our timeline (acres)"),
        ("C", "Farm land needing purchase (acres)"),
        ("D", "Bought from the nations, to date (acres)"),
        ("E", "Bought this year (acres)"),
        ("F", "Annuities paid in goods ($m)"),
        ("G", "Old grants patented this year (acres)"),
        ("H", "Survey and register ($m)"),
        ("I", "Sold to settlers, to date (acres)"),
        ("J", "Instalments received ($m)"),
        ("K", "Grant land paid in and sold ($m)"),
        ("L", "Homeland lease fees ($m)"),
        ("M", "Land Fund net ($m)"),
        ("N", "Commissioners and mounted police ($m)"),
        ("O", "Net cash to the Treasury ($m)"),
        ("P", "Land Fund balance with interest ($m)"),
        ("Q", "Policy balance with interest, incl. police ($m)"),
        ("R", "Assisted passage ($m)"),
        ("S", "Land Fund net / assisted passage, to date"),
    ]
    for col, h in heads:
        ws[f"{col}5"] = h
        ws[f"{col}5"].font = HEAD
        ws[f"{col}5"].fill = PatternFill("solid", fgColor="FF1F4E6B")
        ws[f"{col}5"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[col].width = 14
    ws.row_dimensions[5].height = 54
    ws["A4"] = "Inputs"
    ws["B4"] = "Purchases"
    ws["I4"] = "Receipts"
    ws["M4"] = "Results"
    for c in ("A4", "B4", "I4", "M4"):
        ws[c].font = BOLD
    note(ws["B5"], FARMLAND_SOURCE)
    ws.freeze_panes = "B6"
    p1877 = "INDEX(Economy!$H$6:$H$92,1877-1845)"
    for r in range(FIRST_ROW, LAST_ROW + 1):
        y = 1840 + r
        prev = r - 1
        ws[f"A{r}"] = y
        ws[f"A{r}"].font = BOLD
        ws[f"B{r}"] = round(farmland(y))
        ws[f"B{r}"].font = BLUE
        ws[f"C{r}"] = f"=MAX(0,B{r}*lfund_farm_scale-lfund_grants)"
        # cumulative purchases: the need lfund_lead years ahead, never more than the land outside homelands and grants
        ahead = f"INDEX($C$6:$C$92,MIN(87,{r - 5}+lfund_lead))"
        if r == FIRST_ROW:
            ws[f"D{r}"] = f"=IF(AND(lfund_on=1,$A{r}>=lfund_start),MIN(lfund_area-lfund_home-lfund_grants,{ahead}),0)"
            ws[f"E{r}"] = f"=D{r}"
            ws[f"I{r}"] = f"=IF(AND(lfund_on=1,$A{r}>=lfund_start),MIN(D{r},C{r}),0)"
        else:
            ws[f"D{r}"] = (f"=IF(AND(lfund_on=1,$A{r}>=lfund_start),MIN(lfund_area-lfund_home-lfund_grants,MAX(D{prev},{ahead})),0)")
            ws[f"E{r}"] = f"=D{r}-D{prev}"
            ws[f"I{r}"] = f"=IF(AND(lfund_on=1,$A{r}>=lfund_start),MIN(D{r},MAX(I{prev},C{r})),0)"
        back = lambda col, n: f"IF({r}-{n}<6,0,INDEX(${col}$6:${col}$92,{r}-5-{n}))"
        ws[f"F{r}"] = f"=lfund_price/lfund_term*(D{r}-{back('D', 'lfund_term')})/1000000"
        ws[f"G{r}"] = (f"=IF(AND(lfund_on=1,$A{r}>=lfund_start+1,$A{r}<lfund_start+1+lfund_grant_years),"
                       f"lfund_grants/lfund_grant_years,0)")
        ws[f"H{r}"] = f"=(E{r}+G{r})*lfund_survey/1000000"
        ws[f"J{r}"] = f"=lfund_sale/lfund_inst*(I{r}-{back('I', 'lfund_inst')})*(1-lfund_default)/1000000"
        ws[f"K{r}"] = f"=G{r}*lfund_grant_fee*lfund_sale/1000000"
        ws[f"L{r}"] = (f"=IF(AND(lfund_on=1,$A{r}>=1860),lfund_home*lfund_lease_1932*($A{r}-1859)/(1932-1859)"
                       f"*lfund_rent*lfund_fee,0)/1000000")
        ws[f"M{r}"] = f"=J{r}+K{r}+L{r}-F{r}-H{r}"
        ws[f"N{r}"] = (f"=IF(lfund_on=1,IF(AND($A{r}>=1852,$A{r}<=1870),lfund_police,IF(AND($A{r}>1870,$A{r}<=1890),"
                       f"lfund_police_late,0))*lfund_police_cost*Economy!H{r}/{p1877},0)/1000000")
        ws[f"O{r}"] = f"=M{r}-N{r}"
        if r == FIRST_ROW:
            ws[f"P{r}"] = f"=M{r}"
            ws[f"Q{r}"] = f"=O{r}"
            ws[f"S{r}"] = "=0"
        else:
            ws[f"P{r}"] = f"=P{prev}*(1+lfund_rate)+M{r}"
            ws[f"Q{r}"] = f"=Q{prev}*(1+lfund_rate)+O{r}"
            ws[f"S{r}"] = f"=IF(SUM(R$6:R{r})>0,SUM(M$6:M{r})/SUM(R$6:R{r}),0)"
        ws[f"R{r}"] = f"=People!T{r}*assist_cost*Economy!H{r}/1000000"
        for col in "BCDEGI":
            ws[f"{col}{r}"].number_format = "#,##0;(#,##0);-"
        for col in "FHJKLMNOPQR":
            ws[f"{col}{r}"].number_format = "#,##0.000;(#,##0.000);-"
        ws[f"S{r}"].number_format = "0%;(0%);-"
        for col in "CDEFGHIJKLMNOPQS":
            ws[f"{col}{r}"].font = BLACK

    # summary block
    ws["U5"] = "Summary"
    ws["U5"].font = BOLD
    rows = [
        ("Acres bought, 1851-1932", "=D92", "#,##0"),
        ("Annuities paid ($m)", "=SUM(F6:F92)", "#,##0.0"),
        ("Instalments and grant land ($m)", "=SUM(J6:K92)", "#,##0.0"),
        ("Lease fees ($m)", "=SUM(L6:L92)", "#,##0.0"),
        ("Survey and register ($m)", "=SUM(H6:H92)", "#,##0.0"),
        ("Land Fund net, undiscounted ($m)", "=SUM(M6:M92)", "#,##0.0"),
        ("Police and commissioners ($m)", "=SUM(N6:N92)", "#,##0.0"),
        ("Policy net, undiscounted ($m)", "=SUM(O6:O92)", "#,##0.0"),
        ("Land Fund net, value in 1850 at lfund_rate ($m)", "=NPV(lfund_rate,M11:M92)", "#,##0.0"),
        ("Policy net, value in 1850 at lfund_rate ($m)", "=NPV(lfund_rate,O11:O92)", "#,##0.0"),
        ("Lowest Land Fund balance ($m)", "=MIN(P6:P92)", "#,##0.0"),
        ("First year the Land Fund is back in credit, with interest", '=IF(lfund_on=0,"off",IFERROR(INDEX(A6:A92,MATCH(1,INDEX((P6:P92>=0)*(A6:A92>lfund_start+5),0),0)),"never"))', "0"),
        ("First year the policy is back in credit, with interest", '=IF(lfund_on=0,"off",IFERROR(INDEX(A6:A92,MATCH(1,INDEX((Q6:Q92>=0)*(A6:A92>lfund_start+5),0),0)),"never"))', "0"),
        ("Land Fund net / assisted passage, 1852-1932", "=S92", "0%"),
        ("Our timeline: California's militia claims, 1850s ($m)", 1.293, "#,##0.000"),
    ]
    for i, (label, f, fmt) in enumerate(rows):
        r = 6 + i
        ws[f"U{r}"] = label
        ws[f"U{r}"].font = Font(name="Arial", size=10)
        ws[f"V{r}"] = f
        ws[f"V{r}"].number_format = fmt
        ws[f"V{r}"].font = BLUE if isinstance(f, float) else BLACK
    note(ws["V20"], "Claims submitted to the State Comptroller for expeditions against the Indians: $1,293,179.20 "
                    "(California State Military Museum, 'California and the Indian Wars'). Our-timeline reference only.")
    ws.column_dimensions["U"].width = 44
    ws.column_dimensions["V"].width = 14

    # Budget G: Land Fund net cash
    B = wb["Budget"]
    for r in range(FIRST_ROW, LAST_ROW + 1):
        f = B[f"G{r}"].value
        if "Land Fund" not in f:
            B[f"G{r}"].value = f + f"+'Land Fund'!O{r}"
    B["G5"].value = "Other receipts (Events, Land Fund)"
    note(B["G5"], "From 27 Oct 2026 includes the Land Fund's net cash (Land Fund column O): land sales and fees less "
                  "annuities, survey, commissioners and police (CR-033 draft). Negative when the policy costs more than it brings in.")


# ---------------------------------------------------------------------------

def log_change(wb) -> None:
    C = wb["Changes"]
    r = C.max_row + 1
    C.cell(r, 1).value = "27 Oct: CR-033"
    C.cell(r, 2).value = "Population and the Land Fund"
    C.cell(r, 3).value = ("Population change list P1-P6 (population/workbook_change_list.md): Paths C from the population "
                          "model; DEM-01 early arrivals (dem01_mult); extra arrivals split by origin into People N-P (xo_*), "
                          "Q and R folded in; memo columns for Native Hawaiians and Native Californians; two checks. "
                          "Land Fund sheet (CR-033 draft) feeding Budget G. Events table widened to row 400.")
    for c in range(1, 4):
        C.cell(r, c)._style = copy.copy(C.cell(r - 1, c)._style)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("dst")
    ap.add_argument("--rates")
    ap.add_argument("--no-land-fund", action="store_true")
    a = ap.parse_args()
    rates = None
    if a.rates:
        with open(a.rates, encoding="utf-8") as f:
            rates = {int(r["year"]): float(r["rate"]) for r in csv.DictReader(f)}
    wb = openpyxl.load_workbook(a.src)
    n = expand_events(wb)
    apply_population(wb, rates)
    if not a.no_land_fund:
        apply_land_fund(wb)
    log_change(wb)
    wb.save(a.dst)
    print(f"Events ranges widened in {n} formulas; saved {a.dst}")


if __name__ == "__main__":
    main()
