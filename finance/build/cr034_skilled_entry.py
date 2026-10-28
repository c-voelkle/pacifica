"""Apply CR-034 A24 to the financial workbook: the skilled-entry programme
starts in 1856 instead of 1880.

    python3 finance/build/cr034_skilled_entry.py IN.xlsx OUT.xlsx [--start 1856] [--dem01 X]

IN is the CR-033 workbook (built by cr033_population_landfund.py). The places
cap ("Skilled-entry places / residents") ramps the programme with the
population, so no separate ramp is added. `--dem01` resets the DEM-01
multiplier on 1848-1860 arrivals so that residents at the start of 1861 stay
at about 1.2 million (author decision of 27 Oct 2026): the early skilled
arrivals are part of DEM-01's explanation, not an addition to it.
The script never recalculates; run LibreOffice (xlsx skill scripts/recalc.py)
afterwards.
"""
from __future__ import annotations

import argparse
import copy

import openpyxl
from openpyxl.comments import Comment

AUTHOR = "CR-034 A24, 28 Oct 2026"


def lever_row(L, label: str) -> int:
    for r in range(1, L.max_row + 1):
        if L.cell(r, 2).value == label:
            return r
    raise SystemExit(f"lever {label!r} not found")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("dst")
    ap.add_argument("--start", type=int, default=1856)
    ap.add_argument("--dem01", type=float)
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.src)
    L = wb["Levers"]
    r = lever_row(L, "First year of the skilled-entry programme")
    for col in (3, 4):
        L.cell(r, col).value = a.start
    L.cell(r, 3).comment = Comment(
        f"CR-034 A24 (author, 27-28 Oct 2026): the Polytechnic's recruitment and citizenship with a chair from "
        f"the 1850s start skilled entry in {a.start}, not 1880. The places cap (share of residents) ramps it.",
        AUTHOR, width=320, height=140)
    if a.dem01 is not None:
        d = lever_row(L, "DEM-01: arrivals from the Americas, Europe and China, 1848-1860 (× sheet path)")
        for col in (3, 4):
            L.cell(d, col).value = round(a.dem01, 3)
        L.cell(d, 3).comment = Comment(
            f"Recalibrated with CR-034 A24 to {a.dem01:.3f} so that residents at the start of 1861 stay at about "
            "1.2 million (DEM-01); the early skilled arrivals are part of that explanation.", AUTHOR, width=320, height=120)
    C = wb["Changes"]
    n = C.max_row + 1
    C.cell(n, 1).value = "28 Oct: CR-034"
    C.cell(n, 2).value = "Skilled entry from the 1850s"
    C.cell(n, 3).value = (f"Levers 'First year of the skilled-entry programme' {a.start} (was 1880), CR-034 A24."
                          + (f" DEM-01 multiplier {a.dem01:.3f} keeps 1861 at about 1.2m." if a.dem01 is not None else ""))
    for c in range(1, 4):
        C.cell(n, c)._style = copy.copy(C.cell(n - 1, c)._style)
    wb.save(a.dst)
    print(f"saved {a.dst}")


if __name__ == "__main__":
    main()
