"""Apply CR-040 Stages 3-4 and CR-041 (the Russian war re-planned, Plan C) to the
financial workbook.

    python3 finance/build/cr041_russian_war.py IN.xlsx OUT.xlsx [--rates paths_c_rates_cr041.csv]

IN is the CR-034 workbook (cr033_population_landfund.py, then cr034_skilled_entry.py).
The script never recalculates; run LibreOffice (xlsx skill scripts/recalc.py) afterwards.

What it changes (every new number is a labelled lever or input with its source in a
cell comment; amounts are MODEL estimates):

1. Siberia sheet: the CR-030 schedule (an Ob front, a Pacifican contingent of 4,000
   to 28,000) is replaced by CR-041's: the line seized in one week from 26 May 1918,
   a 50,000-man Pacifican corps at the Urals, the Russian armies at the Urals and on
   the Volga, the Legion homebound in 1919-20, and from 1923 a guarantee reserve
   instead of a standing front. A new railway section, section 5 (Novonikolayevsk-Omsk-
   Chelyabinsk), is added to the transport test, the train-km, the wagon count and
   the operating cost.
2. Events: the Russian war's other money: credit to Komuch and the Directory, the
   Cossack hosts paid in gold, the Kazan gold paying for Pacifican munitions, the
   Black Sea convoys to Wrangel (goods on credit, escort and shipping in cash), the
   White Sea route, the Baltic squadron, Yudenich and Estonia, the reconstruction
   credit, and the Emba, Grozny-Maikop and Baku oil receipts.
3. War credit and Loans: federal Russia recognises part of the Russian share of the
   Entente credit written off in 1918 (lever; B7 in the planning note is still open).
4. Paths C (optional --rates): natural-increase rates from the population model with
   the Pacifican corps' war deaths, for 1918 onwards.
"""
from __future__ import annotations

import argparse
import copy
import csv

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formula.translate import Translator
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter, column_index_from_string as ci
from openpyxl.workbook.defined_name import DefinedName

AUTHOR = "CR-041 / CR-040 Stages 3-4, 30 Oct 2026"
EVENTS_LAST = 400
BLUE = Font(name="Arial", size=10, color="FF0000FF")
BLACK = Font(name="Arial", size=10, color="FF000000")


def note(cell, text: str, h: int = 160) -> None:
    cell.comment = Comment(text, AUTHOR, width=340, height=h)


def lever_row(L, label: str) -> int:
    for r in range(1, L.max_row + 1):
        if L.cell(r, 2).value == label:
            return r
    raise SystemExit(f"lever {label!r} not found")


def add_levers(wb, rows) -> None:
    """rows: (name, label, story, stress, unit, status, source) or a header string."""
    L = wb["Levers"]
    r = L.max_row + 1
    style_src = {c: L.cell(275, c)._style for c in range(1, 12)}
    head_src = {c: L.cell(286, c)._style for c in range(1, 12)}
    for item in rows:
        if isinstance(item, str):
            for c in range(1, 12):
                L.cell(r, c)._style = copy.copy(head_src[c])
            L.cell(r, 2).value = item
            r += 1
            continue
        name, label, story, stress, unit, status, source = item
        for c in range(1, 12):
            L.cell(r, c)._style = copy.copy(style_src[c])
        L.cell(r, 2).value = label
        L.cell(r, 3).value = story
        L.cell(r, 4).value = stress
        L.cell(r, 5).value = f"=IF(case_sel=1,C{r},D{r})"
        L.cell(r, 6).value = unit
        L.cell(r, 7).value = status
        note(L.cell(r, 3), source)
        if name in wb.defined_names:
            raise SystemExit(f"defined name {name!r} already exists")
        wb.defined_names[name] = DefinedName(name, attr_text=f"Levers!$E${r}")
        r += 1


# ---------------------------------------------------------------------------
# 1. Levers
# ---------------------------------------------------------------------------

LEVERS = [
    "— CR-041 / CR-040 Stages 3-4: the Russian war re-planned (30 Oct 2026) —",
    ("sib_L5", "Siberia: Novonikolayevsk-Omsk-Chelyabinsk (section 5, the road to the Ural gate)", 1425, 1425, "km",
     "Model estimate", "Novonikolayevsk-Omsk 627 km and Omsk-Kurgan-Chelyabinsk about 800 km on the original Siberian "
     "line. CR-041 puts the Pacifican corps and the Russian armies' depots at the Urals, so the transport test now "
     "runs to Chelyabinsk. Beyond it (Samara-Zlatoust line, the Volga) the Russian armies run their own supply."),
    ("sib_net5", "Siberia: Net load per train, section 5", 480, 450, "t", "Model estimate",
     "As section 4 (Irkutsk-Ob). Stress: lighter trains on the steppe grades west of Omsk."),
    ("conv_n19", "Black Sea convoys to Novorossiysk in 1919 (one a month from October)", 3, 3, "convoys", "Proposed",
     "CR-041 §3.3: from October 1919 one convoy a month, San Francisco-Panama-Gibraltar-Constantinople-Novorossiysk."),
    ("conv_n20", "Black Sea convoys in 1920 (to the summer, after Moscow and Petrograd)", 6, 8, "convoys", "Model estimate",
     "Moscow falls in March 1920 and Kronstadt in May; the convoys run to mid-1920, then the reconstruction credit "
     "takes over. Stress: two more convoys."),
    ("conv_t", "Tons carried by one Black Sea convoy", 55000, 60000, "t", "Proposed",
     "CR-041 §3.3: six to eight transports, 50,000-60,000 tons: rifles, guns, shell, oil and petrol, lorries."),
    ("conv_mun", "Share of a convoy that is munitions, guns and lorries (the rest fuel)", 0.6, 0.65, "share",
     "Model estimate", "Oil and petrol are bulky and cheap per ton; the rest is the costly cargo."),
    ("conv_vmun", "Value of munitions, guns and lorries", 350, 400, "$ (1913) a ton", "Model estimate",
     "Same unit value as the Siberia sheet's munitions and stores (sib_c_store)."),
    ("conv_vfuel", "Value of fuel oil and petrol", 20, 25, "$ (1913) a ton", "Model estimate",
     "Fuel oil and petrol at Gulf and California prices before the war, delivered."),
    ("conv_esc", "Escort and shipping for one convoy", 1.0, 1.3, "$m (1913)", "Model estimate",
     "Seven transports for a 70-day round trip at sib_c_ship ($0.45m a year), plus one cruiser and two destroyers "
     "on deployment."),
    ("ws_t", "White Sea route, 1919: tons to Archangel, Kotlas and Glazov", 20000, 20000, "t", "Proposed",
     "CR-041 §3.3: two sailings a summer to Archangel from 1919; Dvina barges to Kotlas, rail to Glazov; ~20,000 tons."),
    ("ws_cost", "White Sea route: freight, barges and rail", 40, 55, "$ (1913) a ton", "Model estimate",
     "Ocean freight to Archangel, Dvina barges and the Kotlas-Vyatka rail. The cargo itself is counted on the "
     "Siberia sheet (it replaces rail tonnage to the Perm wing)."),
    ("balt_ops", "Baltic squadron: Pacifican ships' extra running cost on deployment", 15, 20, "$m (1913) a year",
     "Model estimate", "CR-041 §3.3: two turbine cruisers, six destroyers, four minesweepers, a seaplane tender, an "
     "oiler, a hospital ship, eight motor torpedo boats, two transports and a marine battalion, September 1919 to "
     "the return in summer 1920 (4 months in 1919, 7 in 1920). The fleet's ordinary running cost is on Navy."),
    ("yud_19", "Yudenich's rifles, shells and armoured cars, and Estonia's credit: 1919", 12, 15, "$m (1913)",
     "Proposed", "CR-041 §3.3: the squadron lands Yudenich's rifles, shells and forty armoured cars at Reval and "
     "Narva; Estonia, paid in Pacifican credit, makes no separate peace. A loan to the Russian government."),
    ("yud_20", "Yudenich and Estonia: 1920", 6, 8, "$m (1913)", "Model estimate", "Petrograd falls by land in "
     "March-April 1920."),
    ("kazan_gold", "Kazan gold reserve taken on 7 August 1918", 335, 335, "$m", "Approved",
     "CR-041: some 650 million rubles of the Imperial gold reserve, at $0.515 a ruble; in a Reserve Bank escrow at "
     "Chelyabinsk against the Siberian treaty."),
    ("kazan_paid", "Share of the Kazan gold the Directory spends on Pacifican munitions", 0.6, 0.4, "share",
     "Model estimate", "The Directory builds its army on the Kazan gold (§31). The rest backs its currency and passes "
     "to Wrangel's government. Paid 60% in 1919 and 40% in 1920; it offsets the munitions the Siberia sheet "
     "charges to Pacifica."),
    ("dir_18", "Credit to Komuch and the Directory, Treasury cash: 1918", 30, 40, "$m", "Approved",
     "CR-041 N3: Pacifican credit conditional on the land law; the paymaster at Samara withholds the second tranche "
     "until it is enforced. The Directory inherits it, then Wrangel's government. Amount MODEL."),
    ("dir_19", "Credit to the Directory, Treasury cash: 1919", 50, 70, "$m", "Model estimate",
     "The Directory's ministries are funded only while it sits (N2). Amount MODEL."),
    ("cossack_men", "Cossack hosts paid in Pacifican gold (Dutov, Siberian, Transbaikal)", 20000, 25000, "men",
     "Proposed", "CR-041 order of battle: 15,000-25,000 mounted, paid in Pacifican gold; bought, not commanded."),
    ("cossack_pay", "Pay and keep for a Cossack, a year", 150, 200, "$ (1913)", "Model estimate",
     "Between the Siberian subsidy ($100) and the Legion's pay and keep ($600). Seven months in 1918, all of 1919, "
     "half of 1920."),
    ("reco_total", "Reconstruction credit to Wrangel's government, total syndicated", 1000, 1500, "$m", "Proposed",
     "CR-040 (§32): Pacifica's reconstruction credit reaches Wrangel's government in stages, against results. "
     "Size MODEL: about a tenth of Russia's prewar output."),
    ("reco_pub", "Treasury's own share of the reconstruction credit", 0.4, 0.5, "share", "Model estimate",
     "The rest is placed by the San Francisco syndicates with savers, as with the Entente credit."),
    ("reco_years", "Years over which the reconstruction credit is released (from 1920)", 5, 5, "years",
     "Model estimate", "1920-1924, against results."),
    ("russia_recog", "Share of the 1918 Russian write-off that federal Russia recognises", 0.75, 0.5, "share",
     "Proposed", "CR-040: Wrangel and Krivoshein offer a Russia that will recognise the settlement and pay its debts. "
     "Applied to the Treasury's, the savers' and the foreign subscribers' claims alike, and carried on the Entente "
     "funding path. The order of service (B7 in the Anglo-Pacifican planning note) is still open; the Russian gold "
     "sequestered in San Francisco in February 1918 is part of the settlement."),
    ("russia_recog_yr", "Year federal Russia recognises the debt", 1922, 1924, "year", "Model estimate",
     "After Harbin (1921) and with the 1922 standing alliance."),
    ("oil_emba1", "Emba (Dossor) concession: Crown share and royalty, 1919-1922", 0.25, 0.1, "$m a year", "Approved",
     "CR-041 N14: signed with the Western Alash Orda in 1918, kept through the 1921 transfer of Alash to Russia. "
     "Dossor and Makat produced about 0.25 Mt a year before the war. Amount MODEL."),
    ("oil_emba2", "Emba concession: Crown share and royalty, 1923-1932", 1.0, 0.5, "$m a year", "Model estimate",
     "About 1 Mt a year with Pacifican capital, $3 a ton profit, the Crown's 25% share plus royalty."),
    ("oil_groz", "Grozny and Maikop concessions: Crown share, 1922-1932", 2.0, 1.0, "$m a year", "Approved",
     "CR-041 N14: from Wrangel's government against the reconstruction credit. Grozny made about 1.2 Mt a year "
     "before the war. Amount MODEL."),
    ("oil_baku", "Baku joint concession: Pacifica's operating share, Crown share, 1920-1932", 3.0, 1.5, "$m a year",
     "Approved", "CR-041 N14: joint Anglo-Pacifican concession with Azerbaijan, Pacifica operating, Britain the "
     "strategic say. About 4 Mt a year to the Pacifican side. What London collects in 1927 (B8) is still open "
     "and may change this after 1927. Amount MODEL."),
]


# ---------------------------------------------------------------------------
# 2. The Siberia schedule (CR-041 Plan C)
# ---------------------------------------------------------------------------

# Columns: A label, C days, D what happens, E winter loss, F winter flag,
# H Japanese, I Chinese, J Russian & Siberian armies on the line, K Czechoslovak, L Pacifican (000),
# N Japanese share west of Irkutsk, O Chinese share west, R front share of men west,
# S front munitions kg/man-day, V fortification t/day, W railway materials, Y Amur plan t/day,
# AA Kara Sea t, AB depots t/day, AC refugees (000), AD harvest, CK wagons held, DB flotillas & aircraft $m,
# DQ share of section-4 load going west of the Ob, DS planned pairs section 5.
# Depot filling is paced to the Manchuria-Transbaikal branch: heavy in the summer of 1918, when the Amur barges
# run (twice CR-030's plan), lighter in the autumn, when the corps itself is still arriving; at CR-030's pace the
# branch reaches 93% of its military paths in 1918 Q4.
SCHED = [
    ("1918 26 May-Sep", 128, "The Week: the line seized from 26 May; Kazan taken 7 Aug and held 5-12 Sep; Izhevsk relieved",
     0, 0, 30, 6, 15, 42, 30, 0.35, 0, 0.5, 4, 150, 150, 217.391304347826, 0, 250, 300, 0.9, 0, 3, 0.8, 22),
    ("1918 Q4", 92, "Directory at Ufa; the Omsk coup fails (18 Nov); Perm taken 24 Dec; Vance fills the depots",
     0.15, 1, 60, 8, 45, 42, 50, 0.4, 0, 0.4, 4.5, 150, 200, 0, 0, 150, 300, 0.9, 0, 2, 0.8, 22),
    ("1919 Q1", 90, "Glazov and the Kotelnich raid; a Russian army built at Chelyabinsk and Omsk on the Kazan gold",
     0.3, 1, 70, 8, 80, 40, 50, 0.4, 0, 0.4, 3.5, 100, 300, 0, 0, 200, 300, 0.95, 0, 2.5, 0.8, 24),
    ("1919 Q2", 91, "The Volga battle (25 May-15 Jun); Vance on the Zelenodolsk bridge; the Legion's echelons clear",
     0.08, 0, 70, 10, 140, 38, 50, 0.4, 0, 0.55, 7, 100, 350, 219.78021978022, 0, -400, 400, 0.95, 0, 2.5, 0.8, 24),
    ("1919 Q3", 92, "Summer fill: Kotlas line, Amur barges, the Kara Sea convoy; the Legion goes home",
     0, 0, 70, 10, 150, 30, 50, 0.4, 0, 0.45, 3, 80, 300, 326.086956521739, 15000, 300, 300, 1.0, 1500, 2.5, 0.8, 24),
    ("1919 Q4", 92, "Wrangel takes the southern command; Kappel across the frozen Volga: Nizhny and Sormovo",
     0.15, 1, 70, 10, 150, 20, 50, 0.4, 0, 0.55, 6, 60, 100, 0, 0, 200, 250, 1.0, 1500, 2.5, 0.8, 24),
    ("1920 Q1", 91, "Tula, Serpukhov, Moscow; the Perm wing down the frozen Vyatka",
     0.3, 1, 70, 10, 150, 10, 50, 0.4, 0, 0.55, 5, 50, 0, 0, 0, -500, 200, 1.0, 1000, 2, 0.8, 24),
    ("1920 Q2", 91, "Kotelnich: Lenin taken by Maui; Kronstadt; Nikolayevsk",
     0.08, 0, 65, 10, 120, 3, 45, 0.4, 0, 0.45, 4, 40, 0, 87.912087912088, 0, -200, 150, 1.0, 500, 2, 0.8, 24),
    ("1920 Q3", 92, "The trials; drawdown west of the Urals",
     0, 0, 60, 10, 100, 0, 35, 0.35, 0, 0.35, 2.5, 30, 0, 130.434782608696, 15000, 0, 100, 1.0, 0, 1.5, 0.8, 24),
    ("1920 Q4", 92, "Lenin and Trotsky executed; partisans in the Altai",
     0.15, 1, 55, 10, 80, 0, 30, 0.35, 0, 0.35, 2.5, 30, 0, 0, 0, -50, 100, 1.0, 0, 1.5, 0.8, 24),
    ("1921 Q1", 90, "Partisan war behind the line", 0.3, 1, 50, 10, 70, 0, 25, 0.35, 0, 0.35, 2.5, 30, 0, 0, 0,
     -100, 100, 1.0, 0, 1.5, 0.4, 24),
    ("1921 Q2", 91, "Altai partisan campaign; the Tomsk vote", 0.08, 0, 50, 10, 70, 0, 25, 0.35, 0, 0.45, 4, 30, 0,
     87.912087912088, 0, 0, 100, 1.0, 0, 1.5, 0.4, 24),
    ("1921 Q3", 92, "Altai; gas (§32)", 0, 0, 45, 10, 70, 0, 20, 0.35, 0, 0.45, 4, 20, 0, 130.434782608696, 10000,
     0, 100, 1.0, 0, 1.5, 0.4, 24),
    ("1921 Q4", 92, "Treaty of Harbin; drawdown begins", 0.15, 1, 40, 8, 60, 0, 15, 0.3, 0, 0.35, 2.5, 10, 0, 0, 0,
     -50, 50, 1.0, 0, 1, 0.4, 24),
    ("1922 Q1", 90, "Drawdown", 0.3, 1, 35, 8, 50, 0, 10, 0.2, 0, 0.3, 2, 10, 0, 0, 0, 0, 30, 1.0, 0, 1, 0.3, 24),
    ("1922 Q2", 91, "Drawdown", 0.08, 0, 30, 8, 40, 0, 8, 0.1, 0, 0.3, 2, 10, 0, 0, 0, 0, 30, 1.0, 0, 1, 0.3, 24),
    ("1922 Q3", 92, "Drawdown", 0, 0, 25, 6, 30, 0, 6, 0, 0, 0.3, 2, 10, 0, 0, 0, 0, 20, 1.0, 0, 1, 0.3, 24),
    ("1922 Q4", 92, "Standing alliance", 0.15, 1, 20, 6, 20, 0, 5, 0, 0, 0.3, 2, 10, 0, 0, 0, 0, 20, 1.0, 0, 1, 0.3, 24),
]
# 1923-1932: the guarantee of the federal compact (§32), a smaller claim on the railway in reserve.
GUARANTEE = ("Guarantee of the federal compact (tested at winter conditions)", 0.3, 1, 15, 6, 0, 0, 3, 0, 0, 0.2, 2,
             0, 40, 0, 0, None, 10, None, 0, 0.5, 0.2, 24)

SCHED_COLS = ["A", "C", "D", "E", "F", "H", "I", "J", "K", "L", "N", "O", "R", "S", "V", "W", "Y", "AA", "AB", "AC",
              "AD", "CK", "DB", "DQ", "DS"]

SEC5 = [  # (column, header, row-7 formula with {r} and {p} = previous row)
    ("DQ", "Share of the Irkutsk-Ob load going on west of the Ob", None),
    ("DR", "Load, Ob-Chelyabinsk (t a day)", "=AQ{r}*DQ{r}"),
    ("DS", "Planned pairs a day, section 5", None),
    ("DT", "Pairs achieved, section 5", "=DS$6+(DS{r}-DS$6)*sib_works"),
    ("DU", "Military trains available, section 5",
     "=(DT{r}*G{r}-sib_pass)*(1-sib_civ)*(1-sib_fuel-BM{r}*sib_strike)-AU{r}"),
    ("DV", "Freight trains needed, section 5", "=DR{r}/(sib_net5*(1+sib_decup*BB{r}))"),
    ("DW", "Use of military paths, section 5", "=IF(DV{r}=0,0,IF(DU{r}<=0,9.99,DV{r}/DU{r}))"),
]


def apply_siberia(wb) -> None:
    S = wb["Siberia"]
    rows = list(range(6, 34))
    assert len(SCHED) == 18 and all(len(t) == len(SCHED_COLS) for t in SCHED), [len(t) for t in SCHED]
    assert len(GUARANTEE) == len(SCHED_COLS) - 2
    # headers for section 5, styled like the existing ones
    for col, head, _ in SEC5:
        c = S[f"{col}5"]
        c.value = head
        c._style = copy.copy(S["DP5"]._style)
    S["DQ4"].value = "Section 5: Ob-Chelyabinsk (CR-041)"
    S["DQ4"]._style = copy.copy(S["DC4"]._style)
    note(S["DQ5"], "CR-041: the corps fights at the Urals and the Russian armies' depots are at Chelyabinsk. The share of "
         "the westbound tonnage on section 4 that goes on past Novonikolayevsk (the rest is unloaded between Irkutsk "
         "and the Ob, or goes down the Altai branch). 0.8 while the war is at the Urals and on the Volga; 0.4 in 1921, "
         "when the war is in the Altai.")
    note(S["DS5"], "Double track west of the Ob. CR-041 says thirty pairs a day west of Omsk; the Ob-Omsk stretch is "
         "taken as the limit. The storyline's 'twelve pairs a day on the single track east of Omsk' does not match "
         "this sheet's 22 planned pairs on Irkutsk-Ob (CR-030); the sheet's figures are kept.")
    for i, r in enumerate(rows):
        vals = SCHED[i] if i < len(SCHED) else (str(1905 + i),) + (365,) + GUARANTEE
        for col, v in zip(SCHED_COLS, vals):
            if col == "A" and i >= len(SCHED):
                continue  # keep the year label
            if col == "C" and i >= len(SCHED):
                continue
            if v is None:
                continue
            S[f"{col}{r}"].value = v
        for col, _, f in SEC5:
            cell = S[f"{col}{r}"]
            cell._style = copy.copy(S[f"DP{r}"]._style)
            if f is None:
                cell.font = copy.copy(BLUE)
                continue
            cell.value = f.format(r=r)
            cell.font = copy.copy(BLACK)
        if r >= 24:  # 1923-32: depots held level (the existing formula), no Amur or Kara supply
            S[f"Y{r}"].value = 0
            S[f"AA{r}"].value = 0
        # existing formulas extended to section 5
        S[f"CC{r}"].value = f"=MAX(BX{r}:CB{r},DW{r})"
        S[f"CD{r}"].value = (f'=IF(DW{r}>=MAX(BX{r}:CB{r}),"Ob-Chelyabinsk",CHOOSE(MATCH(CC{r},BX{r}:CB{r},0),'
                             f'"Vladivostok-Harbin","Harbin-Manzhouli","Manzhouli-Kitaisky (branch)","Kitaisky-Irkutsk",'
                             f'"Irkutsk-Ob"))')
        S[f"CE{r}"].value = (f"=MIN(BN{r}-BS{r},BO{r}-BT{r},BP{r}-BU{r},BQ{r}-BV{r},BR{r}-BW{r},DU{r}-DV{r})")
        cf = S[f"CF{r}"].value
        assert cf.endswith("+sib_pass))"), cf
        S[f"CF{r}"].value = (cf[:-1] + f"+sib_L5*(DV{r}+AU{r}+(DT{r}*G{r}-sib_pass)*(sib_civ*sib_civused+(1-sib_civ)"
                             f"*sib_fuel)+sib_pass))")
        S[f"CM{r}"].value = S[f"CM{r}"].value + f"+DR{r}/sib_wload*(2*(sib_L4/2+sib_L5/2)/sib_wagkm)"
        di = S[f"DI{r}"].value
        key = f"+sib_L4*(BW{r}+AU{r}))"
        assert key in di, di
        S[f"DI{r}"].value = di.replace(key, f"+sib_L4*(BW{r}+AU{r})+sib_L5*(DV{r}+AU{r}))")
    # row 6 is the first row: no previous row in DT's baseline (DS$6 is itself)
    S["A2"].value = ("One operational schedule (CR-041, Plan C): troops by period, their daily supply, the tonnage on "
                     "each railway section from the sea to the Ural gate, the train paths, locomotives, wagons, port "
                     "and ships that carry it, and what it costs. The first period runs from 26 May to 30 September "
                     "1918. Whole years from 1923 are the guarantee of the federal compact, tested at January "
                     "conditions.")
    S["J5"].value = "Russian and Siberian armies on the line (000)"
    note(S["J5"], "CR-041: the Siberian Army (5,000 in June 1918, 40,000 by October, 100,000 by spring 1919) and the "
         "Directory's army built at Chelyabinsk and Omsk on the Kazan gold (about 120,000 at the front by April 1919), "
         "then Kappel's eastern army and the Perm wing. Armed partly from Pacifican stores (sib_pst_w) and given the "
         "Pacifican subsidy (sib_c_sib). From 1923 federal Russia's own army: not on this sheet.")
    note(S["H5"], "CR-041: three divisions in May 1918, five by September, seven by 1919 (historically about 70,000). "
         "Japan's field divisions hold east of the Yenisei; its line-guard battalions relieve the Legion section by "
         "section as far as Omsk, hence a share west of Irkutsk. From 1923 the Transbaikal protectorate's garrison.")
    note(S["L5"], "CR-041: a Pacifican Expeditionary Corps of 50,000 (two infantry divisions, a railway brigade of "
         "8,000, a river-flotilla detachment, two air squadrons), shipped across the Pacific in April-May 1918, the "
         "only coalition force west of the Yenisei. From 1923 a railway and guarantee reserve.")
    note(S["I5"], "CR-041: a Chinese brigade of about 6,000 guards the Chinese Eastern Railway; Peking's price is that "
         "Chinese soldiers guard it. (The CR-030 schedule had up to 60,000.)")
    note(S["A6"], "26 May-30 September 1918 (128 days): the Week and the Volga summer in one period.")
    S["A36"].value = "By year: what the model uses (1923-32: the guarantee of the federal compact, by quota)"


# ---------------------------------------------------------------------------
# 3. Events
# ---------------------------------------------------------------------------

def add_event(wb, values: list) -> int:
    E = wb["Events"]
    r = next(r for r in range(6, EVENTS_LAST + 1) if E.cell(r, 1).value in (None, ""))
    last_id = max(int(str(E.cell(x, 1).value)[1:]) for x in range(6, EVENTS_LAST + 1)
                  if str(E.cell(x, 1).value).startswith("E"))
    for c in range(1, 14):
        E.cell(r, c)._style = copy.copy(E.cell(174, c)._style)
    E.cell(r, 1).value = f"E{last_id + 1}"
    for i, v in enumerate(values, start=2):
        E.cell(r, i).value = v
    E.cell(r, 9).value = f'=IF(H{r}="","",H{r}*(C{r}-B{r}+1))'
    return r


def price(y: int) -> str:
    return f"INDEX(Paths!$B$6:$B$92,{y - 1845})"


def apply_events(wb) -> None:
    E = wb["Events"]
    # rename the Siberia-fed rows
    for r in range(6, EVENTS_LAST + 1):
        d = E.cell(r, 4).value
        if not isinstance(d, str):
            continue
        if d.startswith("Siberian campaign 19") and d.endswith(": Pacifican spending"):
            y = d[len("Siberian campaign "):len("Siberian campaign ") + 4]
            E.cell(r, 4).value = f"Russian war {y}: Pacifican spending on the line (Siberia sheet)"
            E.cell(r, 12).value = "§31"
            E.cell(r, 13).value = "Siberia sheet (CR-041 Plan C), built up from the operational schedule"
        elif d.startswith("Standing Siberian front "):
            y = d[len("Standing Siberian front "):len("Standing Siberian front ") + 4]
            E.cell(r, 4).value = f"Guarantee of the federal compact {y} (Pacifica's quota)"
            E.cell(r, 11).value = "Proposed"
            E.cell(r, 12).value = "§32"
            E.cell(r, 13).value = ("Siberia sheet: a railway and guarantee reserve instead of a standing front (CR-041); "
                                   "quota × the schedule's cost")
    conv = "(conv_t*(conv_mun*conv_vmun+(1-conv_mun)*conv_vfuel)/1000000)"
    rows = [
        [1918, 1918, "Credit to Komuch and the Directory (Samara tranches)", "Loan", "Asia", "Russia",
         "=dir_18", None, "No", "Approved", "§31 (CR-041 N3)",
         "Treasury cash; conditional on the land law, the second tranche withheld until it is enforced (lever dir_18). "
         "Booked with the Siberian-front partners' loans (Asia group) so that the 1914-18 Entente reconciliation is "
         "untouched"],
        [1919, 1919, "Credit to the Directory (Omsk and Ufa)", "Loan", "Asia", "Russia", "=dir_19", None, "No",
         "Approved", "§31 (CR-041 N2)", "Treasury cash; no ministry funded unless the Directory sits (lever dir_19)"],
        [1918, 1918, "Cossack hosts paid in gold, 1918 (seven months)", "Subsidy & guarantee", None, "Siberia",
         f"=cossack_men*cossack_pay/1000000*7/12*{price(1918)}", None, "Yes", "Proposed", "§31 (CR-041)",
         "Dutov's, the Siberian and the Transbaikal hosts, bought, not commanded"],
        [1919, 1919, "Cossack hosts paid in gold, 1919", "Subsidy & guarantee", None, "Siberia",
         f"=cossack_men*cossack_pay/1000000*{price(1919)}", None, "Yes", "Proposed", "§31 (CR-041)", "As 1918"],
        [1920, 1920, "Cossack hosts paid in gold, 1920 (half year)", "Subsidy & guarantee", None, "Siberia",
         f"=cossack_men*cossack_pay/1000000*0.5*{price(1920)}", None, "Yes", "Proposed", "§31 (CR-041)", "As 1918"],
        [1919, 1919, "Kazan gold: the Directory pays for Pacifican munitions (1919)", "Receipt", None, "Russia",
         "=kazan_gold*kazan_paid*0.6", None, "Yes", "Approved", "§31 (CR-041)",
         "The army built on the Kazan gold; offsets the munitions charged on the Siberia sheet"],
        [1920, 1920, "Kazan gold: the Directory pays for Pacifican munitions (1920)", "Receipt", None, "Russia",
         "=kazan_gold*kazan_paid*0.4", None, "Yes", "Approved", "§31 (CR-041)", "As 1919"],
        [1919, 1919, "Black Sea convoys to Wrangel, 1919: rifles, guns, shell, oil and lorries", "Loan", "Europe",
         "Russia", f"=conv_n19*{conv}*{price(1919)}", None, "No", "Proposed", "§31 (CR-041 §3.3)",
         "Unloaded only when the paymaster certifies a result; a loan to the Russian government"],
        [1920, 1920, "Black Sea convoys to Wrangel, 1920", "Loan", "Europe", "Russia",
         f"=conv_n20*{conv}*{price(1920)}", None, "No", "Proposed", "§31 (CR-041 §3.3)", "As 1919"],
        [1919, 1919, "Black Sea convoys, 1919: escort and shipping", "War & operations", None, "Russia",
         f"=conv_n19*conv_esc*{price(1919)}", None, "No", "Proposed", "§31 (CR-041 §3.3)",
         "Seven transports, one cruiser and two destroyers a convoy"],
        [1920, 1920, "Black Sea convoys, 1920: escort and shipping", "War & operations", None, "Russia",
         f"=conv_n20*conv_esc*{price(1920)}", None, "No", "Proposed", "§31 (CR-041 §3.3)", "As 1919"],
        [1919, 1919, "White Sea route: Archangel, Dvina barges, Kotlas-Glazov", "War & operations", None, "Russia",
         f"=ws_t*ws_cost/1000000*{price(1919)}", None, "No", "Proposed", "§31 (CR-041 §3.3)",
         "Freight only; the cargo is on the Siberia sheet"],
        [1919, 1919, "Baltic squadron on deployment, 1919 (from September)", "War & operations", None, "Russia",
         f"=balt_ops*4/12*{price(1919)}", None, "No", "Approved", "§31 (CR-040 K1, CR-041 §3.3)",
         "Pacifican ships' extra running cost; Japanese ships at Japan's cost"],
        [1920, 1920, "Baltic squadron on deployment, 1920 (Kronstadt and the return)", "War & operations", None,
         "Russia", f"=balt_ops*7/12*{price(1920)}", None, "No", "Approved", "§31 (CR-040 K1)", "As 1919"],
        [1919, 1919, "Yudenich's rifles, shells and armoured cars; Estonia's credit", "Loan", "Europe", "Russia",
         f"=yud_19*{price(1919)}", None, "No", "Proposed", "§31 (CR-041 §3.3)", "Lever yud_19 at 1913 prices"],
        [1920, 1920, "Yudenich and Estonia, 1920", "Loan", "Europe", "Russia", f"=yud_20*{price(1920)}", None, "No",
         "Proposed", "§31 (CR-041 §3.3)", "Lever yud_20 at 1913 prices"],
        [1920, 1924, "Reconstruction credit to Wrangel's government (Treasury share)", "Loan", "Europe", "Russia",
         "=reco_total*reco_pub/reco_years", None, "No", "Proposed", "§32 (CR-040)",
         "In stages, against results; the syndicated rest is placed with savers (not modelled as Treasury cash)"],
        [1919, 1922, "Emba (Dossor) oil concession: Crown share and royalty", "Receipt", None, "Russia",
         "=oil_emba1", None, "No", "Approved", "§32 (CR-041 N14)", "Signed with the Western Alash Orda in 1918"],
        [1923, 1932, "Emba oil concession: Crown share and royalty", "Receipt", None, "Russia", "=oil_emba2", None,
         "No", "Approved", "§32 (CR-041 N14)", "Kept through the 1921 transfer of Alash to Russia"],
        [1922, 1932, "Grozny and Maikop concessions: Crown share", "Receipt", None, "Russia", "=oil_groz", None, "No",
         "Approved", "§32 (CR-041 N14)", "From Wrangel's government against the reconstruction credit"],
        [1920, 1932, "Baku joint concession: Pacifica's operating share, Crown share", "Receipt", None, "Azerbaijan",
         "=oil_baku", None, "No", "Approved", "§32 (CR-041 N14)",
         "Joint with Britain; London's 1927 price (B8) is open and may change this after 1927"],
    ]
    for v in rows:
        add_event(wb, v)


# ---------------------------------------------------------------------------
# 4. Federal Russia recognises the debt
# ---------------------------------------------------------------------------

def apply_recognition(wb) -> None:
    W = wb["War credit"]
    for r in range(6, 93):
        for col, keep in (("N", "N"), ("S", "S"), ("X", "X")):
            f = W[f"{col}{r}"].value
            assert f.startswith(f"=IF(A{r}=1918,"), f
            W[f"{col}{r}"].value = f + f"-IF(A{r}=russia_recog_yr,russia_recog*${keep}$78,0)"
    note(W["N5"], "CR-041 build (30 Oct 2026): a negative loss in the recognition year restores russia_recog of the "
         "1918 Russian write-off (lever russia_recog_yr).")
    Ln = wb["Loans"]
    for r in range(6, 93):
        f = Ln[f"N{r}"].value
        if f == "=0":      # the first year has no opening balance
            continue
        assert f.startswith(f"=IF($A{r}=1918,"), f
        Ln[f"N{r}"].value = f + f"-IF($A{r}=russia_recog_yr,russia_recog*$N$78,0)"


# ---------------------------------------------------------------------------
# 5. Paths C and the change log
# ---------------------------------------------------------------------------

def apply_rates(wb, path: str) -> int:
    P = wb["Paths"]
    n = 0
    with open(path, newline="") as fh:
        for row in csv.DictReader(fh):
            y = int(row["year"])
            if y < 1918:
                continue
            P.cell(y - 1840, 3).value = round(float(row["rate"]), 6)
            n += 1
    return n


def log_change(wb, rates: bool) -> None:
    C = wb["Changes"]
    n = C.max_row + 1
    C.cell(n, 1).value = "30 Oct: CR-041"
    C.cell(n, 2).value = "The Russian war re-planned"
    C.cell(n, 3).value = (
        "CR-041 Plan C and CR-040 Stages 3-4: the Siberia schedule rebuilt (50,000-man corps at the Urals, Russian "
        "armies on the line, a new railway section (section 5) to Chelyabinsk, a guarantee reserve from 1923 instead of a "
        "standing front); new Events for the Directory's credit, Cossack pay, the Kazan gold, the Black Sea, White Sea "
        "and Baltic operations, Yudenich and Estonia, the reconstruction credit and the oil concessions; federal Russia "
        "recognises part of the 1918 write-off." + (" Paths C from 1918 includes the corps' war deaths." if rates else ""))
    for c in range(1, 4):
        C.cell(n, c)._style = copy.copy(C.cell(n - 1, c)._style)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("dst")
    ap.add_argument("--rates")
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.src)
    L = wb["Levers"]
    pw = lever_row(L, "Siberia: Share of Pacifican troops west of Irkutsk")
    for col in (3, 4):
        L.cell(pw, col).value = 0.9
    note(L.cell(pw, 3), "CR-041: the corps is the only coalition force west of the Yenisei; only part of the railway "
         "brigade works east of Irkutsk. Was 0.85 (CR-030).")
    add_levers(wb, LEVERS)
    apply_siberia(wb)
    apply_events(wb)
    apply_recognition(wb)
    n = apply_rates(wb, a.rates) if a.rates else 0
    log_change(wb, bool(a.rates))
    wb.save(a.dst)
    print(f"saved {a.dst}" + (f"; Paths C set for {n} years from 1918" if n else ""))


if __name__ == "__main__":
    main()
