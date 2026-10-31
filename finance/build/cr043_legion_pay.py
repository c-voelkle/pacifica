"""Apply CR-043 P9 (France pays the Czechoslovak Legion) to the financial workbook.

    python3 finance/build/cr043_legion_pay.py IN.xlsx OUT.xlsx

IN is the CR-042 workbook (cr042_huludao.py). The script never recalculates;
run LibreOffice (xlsx skill scripts/recalc.py) afterwards.

The Legion's pay moves to France, which pays it in francs from the Entente credit
Pacifica raised; that credit is already in the Europe-group loans. Pacifica keeps
the cost of feeding and carrying the legionaries and a supplement on the line. The
Siberia lever sib_c_cz (Legion pay and keep, $600 a year at 1913 prices) becomes that
residual cost: $200 (Story) and $250 (Stress), MODEL.
"""
from __future__ import annotations

import argparse
import copy

import openpyxl
from openpyxl.comments import Comment

STORY, STRESS = 200, 250


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("dst")
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.src)
    ref = wb.defined_names["sib_c_cz"].attr_text  # Levers!$E$255
    row = int(ref.split("$")[-1])
    L = wb["Levers"]
    assert L.cell(row, 3).value == 600 and L.cell(row, 4).value == 600, "unexpected sib_c_cz values"
    L.cell(row, 2).value = "Siberia: Czechoslovak legionary, Pacifican keep and line supplement (France pays)"
    L.cell(row, 3).value = STORY
    L.cell(row, 4).value = STRESS
    L.cell(row, 3).comment = Comment(
        "CR-043 P9: France pays the Legion in francs from the Entente credit Pacifica raised "
        "(already in the Europe-group loans). Pacifica feeds it, carries it and pays a supplement "
        "on the line. Was $600 pay and keep. MODEL.", "CR-043, 31 Oct 2026")
    C = wb["Changes"]
    n = C.max_row + 1
    C.cell(n, 1).value = "31 Oct: CR-043"
    C.cell(n, 2).value = "Legion pay"
    C.cell(n, 3).value = ("CR-043 P9: France pays the Czechoslovak Legion; sib_c_cz falls from $600 to $200 (Story) "
                          "and $250 (Stress) a legionary-year, Pacifica's keep and line supplement only.")
    for c in range(1, 4):
        C.cell(n, c)._style = copy.copy(C.cell(n - 1, c)._style)
    wb.save(a.dst)
    print(f"saved {a.dst}")


if __name__ == "__main__":
    main()
