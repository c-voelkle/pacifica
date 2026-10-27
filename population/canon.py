"""Canon run: the People sheet's perimeter and migration, our-timeline vital rates.

What it does
  * Reads the supplied workbook (values only; the file is never written).
  * Starts from the sheet's 1846 Californian residents, adds territories in the
    sheet's join years, and adds the sheet's migration streams year by year,
    each with an age-sex profile from `data/assumptions.toml`.
  * Applies vital rates by census group: our-timeline rates, so the gap to the
    sheet shows how much the canon's divergences must explain.

What it reports (`population/outputs/`)
  * canon_by_year.csv: totals by group, the sheet's total, natural increase,
    men 18-45, women 15-49, children 5-14, ages 15-64, sailors and labour force.
  * reconciliation.md: model against the sheet, the gap by group, and the
    canon checks (Native Hawaiians, Native Californians, DEM-01, crews, labour).

Run:  python3 population/canon.py
"""
from __future__ import annotations

import csv
import math
import sys
from pathlib import Path

import openpyxl

import cohort
import historical
from cohort import Group, Rates, load

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "finance" / "Pacifica_financial_model_1846-1932.xlsx"
OUT = Path(__file__).resolve().parent / "outputs"
FIRST, LAST = 1846, 1932
HAWAII_JOIN = 1856  # the sheet's year (Confederation); Hawaii is de facto core from then


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

def read_sheet() -> dict[int, dict]:
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    cols = {
        "People": {"B": "start", "C": "join", "D": "natural", "E": "am_eu", "F": "chinese", "G": "japanese",
                   "H": "filipino", "I": "leave", "J": "extra", "K": "net", "L": "end", "Q": "asian_guest",
                   "R": "asian_talent", "T": "assisted"},
        "Paths": {"C": "ni_rate", "I": "return_am_eu", "J": "assisted_share"},
        "Navy": {"M": "sailors"},
        "Labour": {"B": "labour_force"},
    }
    data: dict[int, dict] = {}
    for sheet, mapping in cols.items():
        ws = wb[sheet]
        idx = {openpyxl.utils.column_index_from_string(c) - 1: k for c, k in mapping.items()}
        for row in ws.iter_rows(min_row=6, values_only=True):
            if not row or not isinstance(row[0], (int, float)):
                continue
            y = int(row[0])
            if FIRST <= y <= LAST:
                d = data.setdefault(y, {})
                for i, k in idx.items():
                    d[k] = float(row[i] or 0.0)
    levers = wb["Levers"]
    names = {}
    for row in levers.iter_rows(min_row=1, values_only=True):
        if row and isinstance(row[1], str):
            names[row[1]] = row[4]
    data["levers"] = {"dep_asia": float(names.get("Return migration, Asian streams", 0.28)),
                      "family_mult": float(names.get("People per skilled worker admitted (with family)", 2.2))}
    return data


# ---------------------------------------------------------------------------
# Migration by group
# ---------------------------------------------------------------------------

def male_share(stream: str, year: int) -> float:
    return historical.male_share(stream, year, table="migration")


STREAMS = ("americas_europe", "chinese", "japanese", "filipino", "extra")


def migration_profiles(y: int, d: dict, prev: dict | None, lev: dict,
                       streams: tuple[str, ...] = STREAMS) -> dict[str, tuple[list[float], list[float]]]:
    a = load("assumptions")["migration"]
    dep = lev["dep_asia"]
    assisted = d["assisted_share"]
    prof: dict[str, list] = {}

    def stream(name: str, gross: float, leave_share: float, group: str) -> None:
        if gross <= 0 or name not in streams:
            return
        net_unassisted = gross * (1 - assisted) - gross * leave_share
        net_assisted = gross * assisted
        parts = [cohort.migrant_profile(net_unassisted, male_share(name, y),
                                        a["child_share_unassisted"] if net_unassisted > 0 else 0.0),
                 cohort.migrant_profile(net_assisted, 1 - a["assisted_female_share"], a["child_share_family"])]
        prof.setdefault(group, []).extend(parts)

    stream("americas_europe", d["am_eu"], d["return_am_eu"], "white")
    stream("chinese", d["chinese"], dep, "chinese")
    stream("japanese", d["japanese"], dep, "japanese")
    stream("filipino", d["filipino"], dep, "filipino")

    # Extra arrivals drawn by Pacifica's lead: by origin shares (assumptions [extra_origin]),
    # or the sheet's own split (what it adds to its Asian guest-worker and talent stocks).
    eo = load("assumptions")["extra_origin"]
    if eo["mode"] == "shares":
        if "extra" not in streams:
            return {g: cohort.merge(*parts) for g, parts in prof.items()}
        shares = [row for row in eo["periods"] if row[0] <= y][-1][1:]
        fm = lev["family_mult"]
        sk = a["skilled_male_share"]
        for g, share in zip(("white", "japanese", "chinese", "filipino"), shares):
            n = d["extra"] * share
            if n:
                workers = n / fm
                prof.setdefault(g, []).extend([cohort.migrant_profile(workers, sk, 0.0),
                                               cohort.migrant_profile(n - workers, 1 - sk, 0.5)])
        return {g: cohort.merge(*parts) for g, parts in prof.items()}
    asian = 0.0
    if prev:
        for k in ("asian_guest", "asian_talent"):
            asian += d[k] - prev[k] * (1 + d["ni_rate"])
    asian = max(0.0, min(asian, d["extra"]))
    other = d["extra"] - asian
    fm = lev["family_mult"]
    sk = a["skilled_male_share"]

    def skilled(n: float) -> list:
        workers = n / fm
        family = n - workers
        return [cohort.migrant_profile(workers, sk, 0.0), cohort.migrant_profile(family, 1 - sk, 0.5)]

    if "extra" not in streams:
        return {g: cohort.merge(*parts) for g, parts in prof.items()}
    prof.setdefault("white", []).extend(skilled(other))
    weights = {g: d[k] for g, k in (("chinese", "chinese"), ("japanese", "japanese"), ("filipino", "filipino"))}
    tot = sum(weights.values())
    for g, w in weights.items():
        share = w / tot if tot > 0 else (1.0 if g == "chinese" else 0.0)
        if asian * share:
            prof.setdefault(g, []).extend(skilled(asian * share))
    return {g: cohort.merge(*parts) for g, parts in prof.items()}


# ---------------------------------------------------------------------------
# Rates by group
# ---------------------------------------------------------------------------

def make_rates(nh: Rates, native_ca_bound: str):
    nc = load("rates")["native_californian"]

    def rates(group: str, y: int) -> Rates:
        if group == "native_hawaiian":
            return nh
        if group == "native_californian":
            if native_ca_bound == "hawaiian_analogue":
                return nh
            return Rates(nc["floor_e0"], nc["tfr"])  # our_timeline (rescaled to estimates) and sheet_rate (rescaled to the sheet's rate)
        if group == "black":
            return cohort.series_rates("black", y)
        return cohort.series_rates("white", y)
    return rates


def native_ca_target(y: int) -> float | None:
    """Our-timeline bound: follow the estimates (log-linear), hold after 1900."""
    ca = load("census")["california"]
    ys, vs = ca["native_estimate_years"], ca["native_estimate"]
    if y <= ys[0]:
        return None
    if y >= ys[-1]:
        return vs[-1]
    for (y0, v0), (y1, v1) in zip(zip(ys, vs), zip(ys[1:], vs[1:])):
        if y0 <= y <= y1:
            return math.exp(math.log(v0) + (math.log(v1) - math.log(v0)) * (y - y0) / (y1 - y0))


# ---------------------------------------------------------------------------
# Run
# ---------------------------------------------------------------------------

GROUPS = ["native_californian", "white", "black", "chinese", "japanese", "filipino",
          "native_hawaiian", "joined_territories"]


_CACHE: dict = {}


def native_ca_switch() -> str:
    return load("assumptions")["switches"]["native_californians"]


def run(native_ca_bound: str | None = None, joins: bool = True,
        streams: tuple[str, ...] = STREAMS) -> tuple[list[dict], dict]:
    """One cohort run. `joins` and `streams` switch the sheet's territory joins and
    migration streams on or off; the engine is linear in population, so the
    difference between two runs is exactly the contribution of what was switched."""
    native_ca_bound = native_ca_bound or native_ca_switch()
    if "sheet" not in _CACHE:
        _CACHE["sheet"] = read_sheet()
        _CACHE["cal"] = historical.calibrate_native_hawaiian()["chosen"]
    sheet, cal = _CACHE["sheet"], _CACHE["cal"]
    lev = sheet["levers"]
    nh = Rates(cal["e0"], cal["tfr"])
    rates = make_rates(nh, native_ca_bound)
    persp = load("assumptions")["perimeter"]
    hi = load("census")["hawaii"]
    ca = load("census")["california"]

    pop = {g: Group(g) for g in GROUPS}
    # 1846: the sheet's Californian residents = Native Californians + Californios and settlers
    total_1846 = sheet[FIRST]["start"]
    native_1846 = ca["native_estimate"][0]
    pop["native_californian"].add(*_stable(rates("native_californian", FIRST), native_1846, None))
    pop["white"].add(*_stable(rates("white", FIRST), total_1846 - native_1846, 0.55))

    rows = []
    prev = None
    gap = {g: 0.0 for g in GROUPS}  # sheet-implied group population minus model, with compounding
    for y in range(FIRST, LAST + 1):
        d = sheet[y]
        start = {g: pop[g].total() for g in GROUPS}
        # territory joins (the sheet's column C)
        if d["join"] > 0 and joins:
            if y == HAWAII_JOIN:
                nh_n = cohort.interp(hi["native_hawaiian_years"], hi["native_hawaiian"], y)
                pop["native_hawaiian"].add(*_stable(nh, nh_n, None))
                other = d["join"] - nh_n
                pop["chinese"].add(*cohort.migrant_profile(other * 0.2, 0.95, 0.0))
                pop["white"].add(*_stable(rates("white", y), other * 0.8, 0.65))
            else:
                pop["joined_territories"].add(*_stable(rates("white", y), d["join"], persp["joined_male_share"]))
        joined = {g: pop[g].total() for g in GROUPS}
        migr = migration_profiles(y, d, prev, lev, streams)
        births = deaths = 0.0
        ni_by_group = {}
        for g in GROUPS:
            if pop[g].total() <= 0 and g not in migr:
                ni_by_group[g] = 0.0
                continue
            res = cohort.step(pop[g], rates(g, y), migr.get(g))
            births += res["births"]
            deaths += res["deaths"]
            ni_by_group[g] = res["births"] - res["deaths"]
        excess = 0.0
        g = pop["native_californian"]
        target = None
        if native_ca_bound == "our_timeline":
            target = native_ca_target(y + 1)
        elif native_ca_bound == "sheet_rate":
            target = start["native_californian"] * (1 + d["ni_rate"])
        if target is not None and g.total() > 0:
            excess = g.total() - target          # positive: loss beyond the life table
            g.scale(target / g.total())
        c = d["ni_rate"]
        for g in GROUPS:
            change = ni_by_group[g] - (excess if g == "native_californian" else 0.0)
            gap[g] = gap[g] * (1 + c) + c * start[g] - change
        tot = sum(pop[g].total() for g in GROUPS)
        allg = _sum(pop.values())
        row = {"year": y + 1, "model_total": tot, "sheet_total": d["end"],
               "model_natural_increase": births - deaths - excess, "sheet_natural_increase": d["natural"],
               "native_ca_excess_loss": excess, "gap_by_group": dict(gap),
               "sheet_rate_ni_on_group": {g: joined[g] * d["ni_rate"] for g in GROUPS},
               "model_ni_by_group": ni_by_group,
               "men_18_45": allg.between("M", 18, 45), "women_15_49": allg.between("F", 15, 49),
               "children_5_14": allg.between("F", 5, 14) + allg.between("M", 5, 14),
               "age_15_64": allg.between("F", 15, 64) + allg.between("M", 15, 64),
               "males": sum(allg.M), "sailors": d["sailors"], "labour_force": d["labour_force"]}
        for g in GROUPS:
            row[g] = pop[g].total()
        rows.append(row)
        prev = d
    return rows, {"nh": cal, "sheet": sheet}


def _stable(rates: Rates, total: float, male_share: float | None):
    g = cohort.stable_structure(rates, total, male_share)
    return g.F, g.M


def _sum(groups) -> Group:
    s = Group("all")
    for g in groups:
        s.add(g.F, g.M)
    return s


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------

KEY = [1856, 1861, 1865, 1870, 1880, 1887, 1895, 1900, 1914, 1920, 1932]


def fmt(x: float) -> str:
    return f"{x:,.0f}"


def main() -> int:
    OUT.mkdir(exist_ok=True)
    switch = native_ca_switch()
    rows, meta = run()
    by = {r["year"]: r for r in rows}
    bounds = {opt: {r["year"]: r for r in run(opt)[0]}
              for opt in ("our_timeline", "hawaiian_analogue", "sheet_rate")}

    with open(OUT / "canon_by_year.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        cols = ["year", "model_total", "sheet_total"] + GROUPS + [
            "model_natural_increase", "sheet_natural_increase", "native_ca_excess_loss",
            "men_18_45", "women_15_49", "children_5_14", "age_15_64", "males", "sailors", "labour_force"]
        w.writerow(cols)
        for r in rows:
            w.writerow([r["year"]] + [round(r[c]) if isinstance(r[c], float) else r[c] for c in cols[1:]])

    cum_at = {r["year"]: r["gap_by_group"] for r in rows}

    nh = meta["nh"]
    L = ["# Population model: canon run against the People sheet", "",
         "Generated by `population/canon.py` from the supplied workbook (read only). The run uses the sheet's own "
         "perimeter, territory joins and migration streams, and our-timeline vital rates by census group "
         f"(Native Hawaiians: calibrated e0 {nh['e0']:.1f}, TFR {nh['tfr']:.2f}; see `historical.md`). "
         f"Native Californians: switch `{switch}` (data/assumptions.toml). "
         "Migration and territory are identical to the sheet, so the gap below is only the difference in natural increase: "
         "the sheet's single rate against each group's age structure and our-timeline rates. For the sheet's divergence from "
         "our timeline as a whole, and the migration flows, see `flows.md`.", "",
         "## Totals", "",
         "| Year | Sheet | Model | Gap (sheet − model) | Gap % |", "|---|---|---|---|---|"]
    for y in KEY:
        r = by[y]
        gap = r["sheet_total"] - r["model_total"]
        L.append(f"| {y} | {fmt(r['sheet_total'])} | {fmt(r['model_total'])} | {fmt(gap)} | {100 * gap / r['sheet_total']:.1f}% |")
    L += ["", "The sheet adds natural increase as one rate for everyone (Paths column C). The model computes births and deaths "
          "from each group's age and sex structure. Migration is identical in both by construction.", "",
          "## Where the gap comes from", "",
          "Each group's share of the gap: the people the sheet's flat rate gives that group beyond what its age structure "
          "and our-timeline rates give, including their descendants (the sheet's rate compounds on them too).", "",
          "| Group | 1870 | 1887 | 1900 | 1914 | 1932 |", "|---|---|---|---|---|---|"]
    ky = (1870, 1887, 1900, 1914, 1932)
    for g in GROUPS:
        vals = [cum_at[y][g] for y in ky]
        if any(abs(v) > 500 for v in vals):
            L.append(f"| {g.replace('_', ' ')} | " + " | ".join(fmt(v) for v in vals) + " |")
    resid = [by[y]["sheet_total"] - by[y]["model_total"] - sum(cum_at[y].values()) for y in ky]
    L.append("| other (first-year timing of joins, rounding) | " + " | ".join(fmt(v) for v in resid) + " |")
    L.append("| **total gap** | " + " | ".join(fmt(by[y]["sheet_total"] - by[y]["model_total"]) for y in ky) + " |")
    L += ["", "Positive: the sheet has more people in that group than the model; negative: fewer.", ""]

    # Canon checks
    r1856, r1895 = by[1856], by[1895]
    sheet_nh = {y: by[y]["native_hawaiian"] + cum_at[y]["native_hawaiian"] for y in (1860, 1887, 1895)}
    L += ["## Canon checks", "",
          "### Native Hawaiians", "",
          f"Model, our-timeline rates: {fmt(by[1860]['native_hawaiian'])} in 1860, {fmt(by[1887]['native_hawaiian'])} in 1887, "
          f"{fmt(r1895['native_hawaiian'])} in 1895 (census 1896: 39,504). The sheet gives Hawaiians the same growth as everyone, "
          f"which implies {fmt(sheet_nh[1860])} in 1860, {fmt(sheet_nh[1887])} in 1887 and {fmt(sheet_nh[1895])} in 1895. Canon has the Hawaiian population still dying in the 1880s (§20: Noe's figures) and turning upward "
          "about 1895 (§23: Victoria's measures by Noe's method), so the sheet is inconsistent with canon here before 1895. "
          "The turn after 1895 is a canon divergence factor for the next step.", "",
          "### Native Californians", "",
          "Reading of the sheet (to confirm): its 170,000 Californians in 1846 can only be a total that includes the Native "
          f"population, estimated at about {fmt(load('census')['california']['native_estimate'][0])} in 1848; the model starts "
          "Native Californians at that estimate and everyone else at the remainder. If the sheet meant non-Native residents only, "
          "this section and about a third of the total gap change.", "",
          "Canon leaves Indigenous claims and obligations OPEN. Two labelled bounds, against the sheet, which grows them with everyone:", "",
          "| Year | Our-timeline path | Hawaiian-analogue rates (disease without mass killing) | Sheet's rate |", "|---|---|---|---|"]
    for y in (1856, 1870, 1887, 1900, 1932):
        L.append(f"| {y} | " + " | ".join(fmt(bounds[o][y]["native_californian"])
                                           for o in ("our_timeline", "hawaiian_analogue", "sheet_rate")) + " |")
    L += ["", "Our-timeline path: the estimates (150,000 in 1848, 75,000–100,000 in 1853, 30,000 in 1870, 16,000 in 1900), "
          "held at 16,000 after 1900. Choosing between these, or a third path, is an author decision; the model then carries it.", "",
          "### DEM-01 (war_base_assumptions.md)", "",
          f"DEM-01 tests about 1.2m in 1861 and 1.5m in 1865. Sheet: {fmt(by[1861]['sheet_total'])} and {fmt(by[1865]['sheet_total'])}. "
          f"Model: {fmt(by[1861]['model_total'])} and {fmt(by[1865]['model_total'])}. The sheet itself is below DEM-01 in both "
          "years, and natural increase explains little of that (the model is close to the sheet). DEM-01 and the sheet's "
          "perimeter or migration need reconciling: an author decision.", "",
          "### Crews and men of military age", "",
          "| Year | Men 18–45 (model) | Sailors (Navy sheet) | Share |", "|---|---|---|---|"]
    for y in (1870, 1887, 1900, 1914, 1932):
        r = by[y]
        L.append(f"| {y} | {fmt(r['men_18_45'])} | {fmt(r['sailors'])} | {100 * r['sailors'] / r['men_18_45']:.1f}% |")
    L += ["", "### Labour force", "",
          "| Year | Labour force (sheet) | Ages 15–64 (model) | Implied participation |", "|---|---|---|---|"]
    for y in (1870, 1887, 1900, 1914, 1932):
        r = by[y]
        L.append(f"| {y} | {fmt(r['labour_force'])} | {fmt(r['age_15_64'])} | {100 * r['labour_force'] / r['age_15_64']:.0f}% |")
    L += ["", "The sheet's labour force is a flat share of population; against the model's working-age population it implies "
          "a steady participation rate. Whether that fits the intended employment of women (the canon's women's programme) "
          "is an author question; the model has no sourced participation rates yet.", "",
          "## Limits of this version", "",
          "- Our-timeline vital rates throughout: the canon's health measures, the women's programme, and family migration "
          "change rates and are the next step.",
          "- Chinese, Japanese, Filipino, joined-territory and non-Hawaiian Hawaii residents use the white national series.",
          "- Joined territories are one group with a stable age structure; their census race composition is not yet sourced.",
          "- Migrant age and sex profiles are MODEL (`data/assumptions.toml`).",
          "- The Native Hawaiian calibration is weakly identified (see `historical.md`)."]
    (OUT / "reconciliation.md").write_text("\n".join(L) + "\n", encoding="utf-8")
    print("\n".join(L))
    return 0


if __name__ == "__main__":
    sys.exit(main())
